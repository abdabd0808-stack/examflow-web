import os
import sys
import platform
import json
import threading
import time
import urllib.request
import urllib.parse
import subprocess
import requests
import tkinter as tk
import customtkinter as ctk
from tkinter import messagebox, Menu

# Import av pyspellchecker
try:
    from spellchecker import SpellChecker
    HAS_SPELLCHECK_LIB = True
except ImportError:
    HAS_SPELLCHECK_LIB = False

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from docx import Document
from docx.shared import Pt, RGBColor

ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

# --- FIREBASE SKY-KONFIGURASJON ---
FIREBASE_BASE_URL = "https://exam-flow-bedc2-default-rtdb.europe-west1.firebasedatabase.app"

BLOCKED_WIN = [
    "teams.exe", "ms-teams.exe", "msteams.exe", "msteamshost.exe", "TeamsWebView.exe",
    "chrome.exe", "msedge.exe", "firefox.exe", "brave.exe", "opera.exe", "iexplore.exe",
    "vivaldi.exe", "tor.exe", "safari.exe", "arc.exe", "waterfox.exe",
    "discord.exe", "slack.exe", "whatsapp.exe", "telegram.exe",
    "skype.exe", "zoom.exe", "messenger.exe", "signal.exe"
]

BLOCKED_MAC = [
    "Teams", "Microsoft Teams", "Microsoft Teams (work or school)", "msteams", "MSTeams",
    "Safari", "Google Chrome", "Microsoft Edge", "Firefox", "Brave Browser", "Opera",
    "Vivaldi", "Tor Browser", "Arc",
    "Discord", "Slack", "WhatsApp", "Telegram", "Skype", "zoom.us", "Messenger", "Signal"
]

class AdvancedStudentApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("EXAMFLOW — Sikkert Prøverom")
        self.geometry("1100x750")
        self.minsize(950, 650)

        self.code = ""
        self.student_name = ""
        self.exam_data = {}
        self.is_running = False
        self.is_submitted = False
        self.last_saved_text = ""
        self.ignored_words = set()

        # Initiering av norsk stavekontroll fra lokal nb_NO.dic
        self.spell = None
        self.has_spellcheck = False

        if HAS_SPELLCHECK_LIB:
            try:
                self.spell = SpellChecker(language=None)
                
                # Finn stien til nb_NO.dic (støtter både vanlig kjøring og kompilert PyInstaller .exe)
                if getattr(sys, 'frozen', False):
                    base_dir = sys._MEIPASS
                else:
                    base_dir = os.path.dirname(os.path.abspath(__file__))

                dic_path = os.path.join(base_dir, "nb_NO.dic")

                # Sjekk også i arbeidsmappen dersom filen ligger i samme mappe som kjørbart skript
                if not os.path.exists(dic_path):
                    dic_path = "nb_NO.dic"

                if os.path.exists(dic_path):
                    with open(dic_path, "r", encoding="utf-8", errors="ignore") as f:
                        words = [
                            line.split("/")[0].strip().lower() 
                            for line in f 
                            if line.strip() and not line.startswith("#")
                        ]
                    self.spell.word_frequency.load_words(words)
                    self.has_spellcheck = True
                    print(f"Vellykket: Lastet inn {len(words)} norske ord fra {dic_path}!")
                else:
                    print("Advarsel: Fant ikke nb_NO.dic. Stavekontroll vil bruke fallback-ordliste.")
                    fallback_words = ["dette", "er", "en", "prøve", "oppgave", "tekst", "skole", "elev", "lærer"]
                    self.spell.word_frequency.load_words(fallback_words)
                    self.has_spellcheck = True
            except Exception as e:
                print("Feil ved initiering av stavekontroll:", e)
                self.has_spellcheck = False

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        # ================= INNLOGGINGSSKJERM =================
        self.login_card = ctk.CTkFrame(self, fg_color="#ffffff", corner_radius=20, border_width=1, border_color="#cbd5e1", width=420)
        self.login_card.grid(row=0, column=0)
        self.login_card.grid_propagate(False)

        ctk.CTkLabel(
            self.login_card, 
            text="🌿 EXAMFLOW", 
            font=ctk.CTkFont(size=24, weight="bold"),
            text_color="#0284c7"
        ).pack(pady=(35, 5))

        ctk.CTkLabel(
            self.login_card, 
            text="Sikkert digitalt prøverom med Ordbok", 
            font=ctk.CTkFont(size=12),
            text_color="#64748b"
        ).pack(pady=(0, 25))

        ctk.CTkLabel(self.login_card, text="FULLSTENDIG NAVN", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", padx=40)
        self.ent_name = ctk.CTkEntry(self.login_card, placeholder_text="f.eks. Kari Nordmann", height=42, corner_radius=10, fg_color="#f8fafc", text_color="#0f172a", border_color="#cbd5e1")
        self.ent_name.pack(fill="x", padx=40, pady=(4, 15))

        ctk.CTkLabel(self.login_card, text="PRØVEKODE", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", padx=40)
        self.ent_code = ctk.CTkEntry(self.login_card, placeholder_text="f.eks. EXAM-4821", height=42, corner_radius=10, fg_color="#f8fafc", text_color="#0f172a", border_color="#cbd5e1", font=ctk.CTkFont(weight="bold"))
        self.ent_code.pack(fill="x", padx=40, pady=(4, 25))

        self.btn_start = ctk.CTkButton(
            self.login_card, 
            text="Start Prøve", 
            command=self.start_exam,
            height=45,
            corner_radius=10,
            fg_color="#10b981",
            hover_color="#059669",
            font=ctk.CTkFont(size=14, weight="bold")
        )
        self.btn_start.pack(fill="x", padx=40, pady=(0, 30))

        self.exam_frame = ctk.CTkFrame(self, fg_color="#f1f5f9", corner_radius=0)

    def register_student_join(self):
        now_str = time.strftime("%H:%M:%S")
        try:
            clean_name = self.student_name.replace(".", "_").replace("#", "_").replace("$", "_").replace("[", "_").replace("]", "_")
            join_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/students_joined/{clean_name}.json"
            requests.put(join_url, json=now_str, timeout=5)

            alert_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/alerts.json"
            requests.post(alert_url, json={"time": now_str, "student": self.student_name, "message": "Startet prøven."}, timeout=5)
        except Exception as e:
            print("Kunne ikke registrere oppstart i skyen:", e)

    def start_exam(self):
        self.student_name = self.ent_name.get().strip()
        self.code = self.ent_code.get().strip().upper()

        if not self.student_name or not self.code:
            messagebox.showwarning("Mangler info", "Vennligst oppgi både navn og prøvekode.")
            return

        try:
            exam_url = f"{FIREBASE_BASE_URL}/exams/{self.code}.json"
            res = requests.get(exam_url, timeout=5)
            if res.status_code != 200 or not res.json():
                messagebox.showerror("Feil", f"Fant ikke prøve med kode '{self.code}' i skyen. Sjekk koden.")
                return
            self.exam_data = res.json()
        except Exception as e:
            messagebox.showerror("Tilkoblingsfeil", f"Kunne ikke koble til skyen. Sjekk internettforbindelsen:\n{e}")
            return

        self.register_student_join()

        self.is_running = True
        threading.Thread(target=self.security_guard_loop, daemon=True).start()

        self.login_card.grid_forget()
        self.exam_frame.grid(row=0, column=0, sticky="nsew", padx=15, pady=15)
        self.exam_frame.grid_columnconfigure(0, weight=3)
        self.exam_frame.grid_columnconfigure(1, weight=1)
        self.exam_frame.grid_rowconfigure(2, weight=1)

        # --- TOP BAR ---
        top_bar = ctk.CTkFrame(self.exam_frame, fg_color="#ffffff", height=50, corner_radius=10, border_width=1, border_color="#cbd5e1")
        top_bar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 15))

        ctk.CTkLabel(top_bar, text=f"👤 {self.student_name}  |  Kode: {self.code}", font=ctk.CTkFont(size=12, weight="bold"), text_color="#0284c7").pack(side="left", padx=20)
        self.lbl_word_count = ctk.CTkLabel(top_bar, text="Ord: 0", font=ctk.CTkFont(size=12, weight="bold"), text_color="#10b981")
        self.lbl_word_count.pack(side="right", padx=20)

        # --- VENSTRE SIDE: TEKSTFELT ---
        left_box = ctk.CTkFrame(self.exam_frame, fg_color="transparent")
        left_box.grid(row=1, column=0, rowspan=2, sticky="nsew", padx=(0, 10))
        left_box.grid_columnconfigure(0, weight=1)
        left_box.grid_rowconfigure(1, weight=1)

        prompt_card = ctk.CTkFrame(left_box, fg_color="#ffffff", corner_radius=10, border_width=1, border_color="#cbd5e1")
        prompt_card.grid(row=0, column=0, sticky="ew", pady=(0, 10))

        instruction_text = self.exam_data.get("prompt") or self.exam_data.get("instruction") or self.exam_data.get("oppgave") or "Ingen instruksjon oppgitt"

        ctk.CTkLabel(prompt_card, text="OPPGAVE:", font=ctk.CTkFont(size=10, weight="bold"), text_color="#64748b").pack(anchor="w", padx=15, pady=(10, 2))
        ctk.CTkLabel(prompt_card, text=instruction_text, font=ctk.CTkFont(size=13), text_color="#0f172a", justify="left").pack(anchor="w", padx=15, pady=(0, 10))

        self.txt_answer = ctk.CTkTextbox(
            left_box, 
            font=ctk.CTkFont(size=14), 
            fg_color="#ffffff", 
            text_color="#0f172a", 
            corner_radius=10,
            border_width=1,
            border_color="#cbd5e1"
        )
        self.txt_answer.grid(row=1, column=0, sticky="nsew", pady=(0, 5))

        # Tag for rød understreking
        self.txt_answer._textbox.tag_config("misspelled", underline=True, underlinefg="#ef4444")

        # Blokker lim-inn
        self.txt_answer.bind("<Control-v>", lambda e: "break")
        self.txt_answer.bind("<Control-V>", lambda e: "break")
        self.txt_answer.bind("<Command-v>", lambda e: "break")
        
        self.txt_answer.bind("<KeyRelease>", self.on_text_key_release)
        
        # Høyreklikkmeny
        self.txt_answer._textbox.bind("<Button-3>", self.show_spelling_context_menu)
        self.txt_answer._textbox.bind("<Button-2>", self.show_spelling_context_menu)

        initial_spell_status = "✏️ Norsk stavekontroll: Aktiv (Bokmål)." if self.has_spellcheck else "✏️ Stavekontroll: Deaktivert."
        self.lbl_spell_warning = ctk.CTkLabel(
            left_box, 
            text=initial_spell_status, 
            font=ctk.CTkFont(size=11), 
            text_color="#64748b",
            anchor="w"
        )
        self.lbl_spell_warning.grid(row=2, column=0, sticky="ew", pady=(0, 10))

        btn_action_frame = ctk.CTkFrame(left_box, fg_color="transparent")
        btn_action_frame.grid(row=3, column=0, sticky="ew")
        btn_action_frame.grid_columnconfigure(0, weight=1)

        self.btn_submit = ctk.CTkButton(
            btn_action_frame, 
            text="🚀 Lever Besvarelse", 
            command=self.submit_exam,
            height=45,
            corner_radius=10,
            fg_color="#2563eb",
            hover_color="#1d4ed8",
            font=ctk.CTkFont(size=14, weight="bold")
        )
        self.btn_submit.grid(row=0, column=0, columnspan=3, sticky="ew")

        self.btn_download_pdf = ctk.CTkButton(
            btn_action_frame, 
            text="📄 Last ned PDF", 
            command=self.download_pdf,
            height=45,
            corner_radius=10,
            fg_color="#0284c7",
            hover_color="#0369a1",
            font=ctk.CTkFont(size=13, weight="bold")
        )

        self.btn_download_word = ctk.CTkButton(
            btn_action_frame, 
            text="📝 Last ned Word (.docx)", 
            command=self.download_word,
            height=45,
            corner_radius=10,
            fg_color="#15803d",
            hover_color="#166534",
            font=ctk.CTkFont(size=13, weight="bold")
        )

        # --- HØYRE SIDE: ORDBOK ---
        dict_card = ctk.CTkFrame(self.exam_frame, fg_color="#ffffff", corner_radius=10, border_width=1, border_color="#cbd5e1")
        dict_card.grid(row=1, column=1, rowspan=2, sticky="nsew")
        dict_card.grid_columnconfigure(0, weight=1)
        dict_card.grid_rowconfigure(3, weight=1)

        ctk.CTkLabel(dict_card, text="📖 BOKMÅLSORDBOKA", font=ctk.CTkFont(size=12, weight="bold"), text_color="#0284c7").grid(row=0, column=0, padx=15, pady=(15, 10), sticky="w")

        search_frame = ctk.CTkFrame(dict_card, fg_color="transparent")
        search_frame.grid(row=1, column=0, sticky="ew", padx=15, pady=(0, 10))
        search_frame.grid_columnconfigure(0, weight=1)

        self.ent_dict = ctk.CTkEntry(search_frame, placeholder_text="Søk i ordboka...", height=35, corner_radius=8, fg_color="#f8fafc", text_color="#0f172a", border_color="#cbd5e1")
        self.ent_dict.grid(row=0, column=0, sticky="ew", padx=(0, 5))
        self.ent_dict.bind("<Return>", lambda e: self.search_dict())

        btn_dict_search = ctk.CTkButton(search_frame, text="Søk", width=60, height=35, corner_radius=8, command=self.search_dict, fg_color="#0284c7", hover_color="#0369a1")
        btn_dict_search.grid(row=0, column=1)

        self.txt_dict_result = ctk.CTkTextbox(
            dict_card, 
            font=ctk.CTkFont(size=12), 
            fg_color="#f8fafc", 
            text_color="#0f172a", 
            corner_radius=8,
            border_width=1,
            border_color="#cbd5e1",
            wrap="word"
        )
        self.txt_dict_result.grid(row=3, column=0, sticky="nsew", padx=15, pady=(0, 15))
        self.txt_dict_result.insert("1.0", "Søk etter ord her for forklaringer.")
        self.txt_dict_result.configure(state="disabled")

        threading.Thread(target=self.live_autosave_loop, daemon=True).start()

    def check_live_spelling(self):
        if not getattr(self, 'has_spellcheck', False) or not self.spell:
            return

        textbox = self.txt_answer._textbox
        textbox.tag_remove("misspelled", "1.0", "end")

        full_text = textbox.get("1.0", "end-1c")
        if not full_text.strip():
            return

        lines = full_text.split("\n")
        for line_idx, line in enumerate(lines, start=1):
            start_col = 0
            words = line.split(" ")
            for w in words:
                clean_w = w.strip(".,!?\"'()[]:;")
                if clean_w.isalpha() and clean_w.lower() not in self.ignored_words:
                    if self.spell.unknown([clean_w.lower()]):
                        w_start = line.find(clean_w, start_col)
                        if w_start != -1:
                            w_end = w_start + len(clean_w)
                            start_index = f"{line_idx}.{w_start}"
                            end_index = f"{line_idx}.{w_end}"
                            textbox.tag_add("misspelled", start_index, end_index)
                start_col += len(w) + 1

    def show_spelling_context_menu(self, event):
        if not getattr(self, 'has_spellcheck', False) or not self.spell:
            return

        textbox = self.txt_answer._textbox
        click_index = textbox.index(f"@{event.x},{event.y}")
        
        word_start = textbox.index(f"{click_index} wordstart")
        word_end = textbox.index(f"{click_index} wordend")
        clicked_word = textbox.get(word_start, word_end).strip(".,!?\"'()[]:;")

        if not clicked_word or not clicked_word.isalpha():
            return

        menu = Menu(self, tearoff=0)

        if self.spell.unknown([clicked_word.lower()]) and clicked_word.lower() not in self.ignored_words:
            candidates = list(self.spell.candidates(clicked_word.lower()) or [])[:5]
            
            if candidates:
                for candidate in candidates:
                    formatted_candidate = candidate.capitalize() if clicked_word[0].isupper() else candidate
                    menu.add_command(
                        label=formatted_candidate, 
                        font=("Segoe UI", 10, "bold"),
                        command=lambda c=formatted_candidate, s=word_start, e=word_end: self.replace_word(s, e, c)
                    )
            else:
                menu.add_command(label="(Ingen forslag)", state="disabled")
            
            menu.add_separator()
            menu.add_command(label="Ignorer alle", command=lambda: self.ignore_word(clicked_word))
        else:
            menu.add_command(label="✅ Riktig stavet", state="disabled")

        menu.tk_popup(event.x_root, event.y_root)

    def replace_word(self, start_idx, end_idx, new_word):
        textbox = self.txt_answer._textbox
        textbox.delete(start_idx, end_idx)
        textbox.insert(start_idx, new_word)
        self.check_live_spelling()

    def ignore_word(self, word):
        self.ignored_words.add(word.lower())
        self.check_live_spelling()

    def auto_capitalize(self):
        try:
            cursor_pos = self.txt_answer.index("insert")
            text = self.txt_answer.get("1.0", "end-1c")

            if not text:
                return

            if len(text) == 1 and text[0].islower():
                self.txt_answer.delete("1.0", "end")
                self.txt_answer.insert("1.0", text[0].upper())
                self.txt_answer.mark_set("insert", cursor_pos)
                return

            if len(text) >= 3:
                last_three = text[-3:]
                if last_three[0] in [".", "!", "?"] and last_three[1] == " " and last_three[2].islower():
                    new_text = text[:-1] + text[-1].upper()
                    self.txt_answer.delete("1.0", "end")
                    self.txt_answer.insert("1.0", new_text)
                    self.txt_answer.mark_set("insert", cursor_pos)
        except Exception:
            pass

    def on_text_key_release(self, event=None):
        if event and event.keysym not in ["BackSpace", "Delete", "Left", "Right", "Up", "Down"]:
            self.auto_capitalize()
        self.update_word_count()
        self.check_live_spelling()

    def update_word_count(self, event=None):
        text = self.txt_answer.get("1.0", "end-1c").strip()
        words = len(text.split()) if text else 0
        self.lbl_word_count.configure(text=f"Ord: {words}")

    def live_autosave_loop(self):
        clean_name = self.student_name.replace(".", "_").replace("#", "_").replace("$", "_").replace("[", "_").replace("]", "_")
        base_student_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/live_texts/{clean_name}"
        
        last_history_save_time = 0

        while self.is_running and not self.is_submitted:
            try:
                current_text = self.txt_answer.get("1.0", "end-1c")
                current_word_count = len(current_text.split()) if current_text.strip() else 0
                now_str = time.strftime("%H:%M:%S")
                now_timestamp = int(time.time() * 1000)

                if current_text != self.last_saved_text:
                    payload = {
                        "text": current_text,
                        "words": current_word_count,
                        "status": "Aktiv"
                    }
                    requests.put(f"{base_student_url}.json", json=payload, timeout=3)
                    self.last_saved_text = current_text

                if time.time() - last_history_save_time >= 30:
                    history_point = {
                        "time": now_str,
                        "words": current_word_count,
                        "timestamp": now_timestamp
                    }
                    requests.post(f"{base_student_url}/history.json", json=history_point, timeout=3)
                    last_history_save_time = time.time()

            except Exception:
                pass
            time.sleep(1.0)

    def security_guard_loop(self):
        CREATE_NO_WINDOW = 0x08000000
        is_mac = platform.system() == "Darwin"

        while self.is_running and not self.is_submitted:
            try:
                if is_mac:
                    for proc_name in BLOCKED_MAC:
                        subprocess.run(["pkill", "-9", "-f", proc_name], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                else:
                    for proc_name in BLOCKED_WIN:
                        subprocess.run(["taskkill", "/F", "/T", "/IM", proc_name], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW)
            except Exception:
                pass
            time.sleep(0.1)

    def submit_exam(self):
        answer = self.txt_answer.get("1.0", "end-1c").strip()
        if not answer:
            messagebox.showwarning("Tom besvarelse", "Du kan ikke levere en tom besvarelse!")
            return

        confirm = messagebox.askyesno("Bekreft levering", "Er du sikker på at du vil levere besvarelsen?\nNår du leverer, låses teksten for redigering.")
        if not confirm:
            return

        now_str = time.strftime("%H:%M:%S")

        try:
            clean_name = self.student_name.replace(".", "_").replace("#", "_").replace("$", "_").replace("[", "_").replace("]", "_")

            sub_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/submissions/{clean_name}.json"
            requests.put(sub_url, json=answer, timeout=5)

            time_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/submission_times/{clean_name}.json"
            requests.put(time_url, json=now_str, timeout=5)

            alert_url = f"{FIREBASE_BASE_URL}/exams/{self.code}/alerts.json"
            requests.post(alert_url, json={"time": now_str, "student": self.student_name, "message": "Leverte prøven."}, timeout=5)

        except Exception as e:
            messagebox.showerror("Feil ved lagring i skyen", f"Kunne ikke sende til skyen: {e}")
            return

        self.txt_answer.configure(state="disabled")
        self.is_submitted = True

        pdf_path = self.generate_pdf_file(answer)
        docx_path = self.generate_word_file(answer)

        self.btn_submit.configure(state="disabled", text="✅ Levert", fg_color="#059669")
        
        self.btn_submit.grid(row=0, column=2, sticky="ew", padx=(3, 0))
        self.btn_download_pdf.grid(row=0, column=0, sticky="ew", padx=(0, 3))
        self.btn_download_word.grid(row=0, column=1, sticky="ew", padx=(3, 3))

        messagebox.showinfo("Levering Vellykket!", f"Besvarelsen din er levert kl. {now_str}.\n\nFiler lagret i Nedlastinger:\n• {os.path.basename(pdf_path)}\n• {os.path.basename(docx_path)}")

    def download_pdf(self):
        answer_text = self.txt_answer.get("1.0", "end-1c").strip()
        pdf_path = self.generate_pdf_file(answer_text)
        messagebox.showinfo("PDF Lagret!", f"PDF-en ligger klar i Nedlastinger-mappen:\n\n{os.path.basename(pdf_path)}")

    def download_word(self):
        answer_text = self.txt_answer.get("1.0", "end-1c").strip()
        docx_path = self.generate_word_file(answer_text)
        messagebox.showinfo("Word-dokument Lagret!", f"Word-filen ligger klar i Nedlastinger-mappen:\n\n{os.path.basename(docx_path)}")

    def generate_pdf_file(self, answer_text):
        downloads_folder = os.path.join(os.path.expanduser("~"), "Downloads")
        if not os.path.exists(downloads_folder):
            downloads_folder = os.getcwd()

        filename = f"Levert_Besvarelse_{self.student_name.replace(' ', '_')}_{self.code}.pdf"
        file_path = os.path.join(downloads_folder, filename)

        doc = SimpleDocTemplate(file_path, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor("#0f172a"), spaceAfter=10)
        meta_style = ParagraphStyle('MetaStyle', parent=styles['Normal'], fontSize=10, textColor=colors.HexColor("#475569"), spaceAfter=15)
        body_style = ParagraphStyle('BodyStyle', parent=styles['Normal'], fontSize=11, leading=16, textColor=colors.HexColor("#1e293b"), spaceAfter=12)

        story = [
            Paragraph(f"Offisiell Eksamensbesvarelse — {self.code}", title_style),
            Paragraph(f"<b>Elev:</b> {self.student_name} | <b>Prøvekode:</b> {self.code} | <b>Levert dato:</b> {time.strftime('%d.%m.%Y kl. %H:%M:%S')}", meta_style),
            Spacer(1, 10)
        ]

        prompt_text = self.exam_data.get('prompt') or self.exam_data.get('instruction') or self.exam_data.get('oppgave') or ''
        if prompt_text:
            story.append(Paragraph(f"<b>Oppgave:</b> <i>{prompt_text}</i>", meta_style))
            story.append(Spacer(1, 15))

        paragraphs = answer_text.split('\n')
        for p in paragraphs:
            if p.strip():
                safe_p = p.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(safe_p, body_style))
            else:
                story.append(Spacer(1, 8))

        doc.build(story)
        return file_path

    def generate_word_file(self, answer_text):
        downloads_folder = os.path.join(os.path.expanduser("~"), "Downloads")
        if not os.path.exists(downloads_folder):
            downloads_folder = os.getcwd()

        filename = f"Levert_Besvarelse_{self.student_name.replace(' ', '_')}_{self.code}.docx"
        file_path = os.path.join(downloads_folder, filename)

        doc = Document()
        title = doc.add_heading(level=1)
        run_title = title.add_run(f"Offisiell Eksamensbesvarelse — {self.code}")
        run_title.font.size = Pt(18)
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        meta = doc.add_paragraph()
        p_meta = meta.add_run(f"Elev: {self.student_name} | Prøvekode: {self.code} | Levert dato: {time.strftime('%d.%m.%Y kl. %H:%M:%S')}")
        p_meta.font.size = Pt(10)
        p_meta.font.italic = True
        p_meta.font.color.rgb = RGBColor(71, 85, 105)

        prompt_text = self.exam_data.get('prompt') or self.exam_data.get('instruction') or self.exam_data.get('oppgave') or ''
        if prompt_text:
            prompt_p = doc.add_paragraph()
            prompt_p.add_run("Oppgave: ").bold = True
            prompt_p.add_run(prompt_text).italic = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        for p in answer_text.split('\n'):
            if p.strip():
                doc_p = doc.add_paragraph(p)
                doc_p.paragraph_format.space_after = Pt(8)
                doc_p.paragraph_format.line_spacing = 1.15

        doc.save(file_path)
        return file_path

    def search_dict(self):
        query = self.ent_dict.get().strip()
        if not query:
            return

        self.txt_dict_result.configure(state="normal")
        self.txt_dict_result.delete("1.0", "end")
        self.txt_dict_result.insert("1.0", f"Søker etter '{query}'...\n\n")

        threading.Thread(target=self._fetch_ordbokene_data, args=(query,), daemon=True).start()

    def _fetch_ordbokene_data(self, query):
        try:
            clean_query = query.lower().strip()
            if clean_query.startswith("å "):
                clean_query = clean_query[2:].strip()

            encoded_query = urllib.parse.quote(clean_query)
            search_url = f"https://ord.uib.no/api/suggest?q={encoded_query}&dict=bm&n=8"
            headers = {'User-Agent': 'Mozilla/5.0'}
            
            req = urllib.request.Request(search_url, headers=headers)
            with urllib.request.urlopen(req, timeout=6) as resp:
                search_res = json.loads(resp.read().decode('utf-8'))

            exact_words = search_res.get('a', {}).get('exact', [])
            freetext_words = search_res.get('a', {}).get('freetext', [])
            all_items = exact_words + freetext_words

            if not all_items:
                self.after(0, lambda: self._update_dict_ui(f"FANT INGEN TREFF\n\nKunne ikke finne ordet '{query}'."))
                return

            output_text = f"🔍 RESULTATER FOR '{query.upper()}':\n\n"
            seen_words = set()

            for item in all_items:
                word_name = item[0] if isinstance(item, list) else item
                if not isinstance(word_name, str) or word_name in seen_words:
                    continue
                seen_words.add(word_name)

                output_text += f"📌 {word_name.upper()}\n• Registrert i Bokmålsordboka.\n\n"

            self.after(0, lambda: self._update_dict_ui(output_text))

        except Exception as e:
            self.after(0, lambda: self._update_dict_ui(f"Feil ved oppslag: {e}"))

    def _update_dict_ui(self, text):
        self.txt_dict_result.configure(state="normal")
        self.txt_dict_result.delete("1.0", "end")
        self.txt_dict_result.insert("1.0", text)
        self.txt_dict_result.configure(state="disabled")

if __name__ == "__main__":
    app = AdvancedStudentApp()
    app.mainloop()